"""Layout of Austrian legal documents with python-docx: Schriftsätze and Verträge.

The drafting model writes the text line by line with markers; this module sets the layout, so
every brief and every contract looks the same. Page 1 carries only the Rubrum (brief) or the
title and the parties (contract); the text starts on page 2.

Schriftsatz:
    [ADRESSAT] An das            [AZ] 12 Cg 34/26x
    [ADRESSAT] Landesgericht für ZRS Wien
    [PARTEI] Klagende Partei: | Name; Anschrift; FN
    [PARTEI] vertreten durch: | Kanzlei; Anschrift; Code (Vollmacht erteilt gemäß § 30 Abs 2 ZPO)
    [PARTEI] Beklagte Partei: | …
    [WEGEN] EUR 18.400,-- s.A.   [STREITWERT] EUR 18.400,--
    [RUBRUM] Berufungsinteresse: | EUR 18.400,--      (any further Rubrum line)
    [TITEL] KLAGE                [UNTERTITEL] …            [BEILAGEN] ./A bis ./F
    --- page 2 ---
    [H] I. Sachverhalt           [P] 1. …                  [BEWEIS] ./A Rechnung; PV; Zeuge …
    [ANTRAG] Die beklagte Partei ist schuldig, …
    [SCHLUSS] Wien, am [Datum] | Klagende Partei

Vertrag:
    [VERTRAG] MIETVERTRAG        [UNTERTITEL] über das Geschäftslokal …
    [PARTEI] Name; Anschrift; FN | Vermieterin              (one line per party)
    --- page 2 ---
    [H] 1. Mietgegenstand        [P] 1.1 …
    [UNTERSCHRIFT] Wien, am [Datum] | Name der Partei       (one line per party)
    [ANLAGE] ./1 Lageplan

Lines without a marker continue the previous paragraph. Placeholders in [ … ] are
highlighted yellow.

    import sys; sys.path.insert(0, '/home/user/.aiss')
    from doc_layout import render
    render(TEXT, '/home/user/Klage.docx')     # brief or contract, detected from the markers
"""
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MARKERS = ('ADRESSAT', 'AZ', 'PARTEI', 'WEGEN', 'STREITWERT', 'RUBRUM', 'TITEL', 'UNTERTITEL', 'BEILAGEN', 'H', 'P', 'BEWEIS', 'ANTRAG', 'SCHLUSS', 'VERTRAG', 'UNTERSCHRIFT', 'ANLAGE')
LINE = re.compile(r'^\s*\[(%s)\]\s?(.*)$' % '|'.join(MARKERS))
PLACEHOLDER = re.compile(r'(\[[^\[\]]*\])')
FONT = 'Arial'
LABEL_W, VALUE_W = Cm(4.5), Cm(12.0)  # A4 minus margins (2.5 + 2 cm) = 16.5 cm


def parse(text):
    """[(marker, text)] – unmarked lines join the paragraph before them."""
    items = []
    for raw in str(text).replace('\r\n', '\n').split('\n'):
        m = LINE.match(raw)
        if m:
            items.append([m.group(1), m.group(2).strip()])
        elif raw.strip():
            if items and items[-1][0] in ('P', 'ANTRAG', 'H', 'BEWEIS'):
                items[-1][1] += ' ' + raw.strip()
            else:
                items.append(['P', raw.strip()])
    return items


def runs(par, text, bold=False, italic=False, size=None):
    for part in PLACEHOLDER.split(text):
        if not part:
            continue
        r = par.add_run(part)
        if bold:
            r.bold = True
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
        if PLACEHOLDER.fullmatch(part) and not re.fullmatch(r'\[\d+\]', part):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def fmt(par, before=0, after=6, align=None, left=None, right=None, first=None, spacing=1.15, keep_next=False):
    f = par.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), spacing
    if align is not None:
        par.alignment = align
    if left is not None:
        f.left_indent = Cm(left)
    if right is not None:
        f.right_indent = Cm(right)
    if first is not None:
        f.first_line_indent = Cm(first)
    if keep_next:
        f.keep_with_next = True
    return par


def tab_stop(par, cm):
    par.paragraph_format.tab_stops.add_tab_stop(Cm(cm))


def field(par, instr, size=9):
    def fld(kind):
        r = par.add_run()
        e = OxmlElement('w:fldChar')
        e.set(qn('w:fldCharType'), kind)
        r._r.append(e)
        r.font.size = Pt(size)
    fld('begin')
    r = par.add_run()
    t = OxmlElement('w:instrText')
    t.set(qn('xml:space'), 'preserve')
    t.text = instr
    r._r.append(t)
    fld('separate')
    par.add_run('1').font.size = Pt(size)
    fld('end')


def widths(table, *cols):
    """Column widths in the table grid as well as in every cell (Word and LibreOffice read different ones)."""
    for col, w in zip(table.columns, cols):
        col.width = w
    for row in table.rows:
        for cell, w in zip(row.cells, cols):
            cell.width = w


def borderless(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + side)
        b.set(qn('w:val'), 'nil')
        borders.append(b)
    tbl_pr.append(borders)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False


def cell_text(cell, lines, bold_first=False, align=None, size=None):
    cell.paragraphs[0].text = ''
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        fmt(p, after=0, align=align, spacing=1.15)
        runs(p, line, bold=bold_first and i == 0, size=size)


def setup(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0), Cm(2.5), Cm(2.0)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = FONT, Pt(11)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(k), FONT)
    foot = sec.footer.paragraphs[0]
    fmt(foot, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    foot.add_run('Seite ').font.size = Pt(9)
    field(foot, 'PAGE')
    foot.add_run(' von ').font.size = Pt(9)
    field(foot, 'NUMPAGES')


# ---- page 1 of a brief: the Rubrum is spread over the page instead of sticking to the top ----
LINE_PT = 15.2                                       # one line of Arial 11 at 1.15 spacing
PAGE_PT = (29.7 - 2.5 - 2.0) / 2.54 * 72 - 40        # usable height of page 1 minus footer and safety


def _lines(text, chars):
    return max(1, -(-len(text) // chars))


def spread(adressat, az, rows, items):
    """Extra space (pt) after the addressee, between parties, before wegen/Streitwert and before the title."""
    h = LINE_PT * max(len(adressat) or 3, len(az)) + 10
    party_gaps = value_gaps = 0
    for i, (kind, label, value) in enumerate(rows):
        vals = [x.strip() for x in value.split(';') if x.strip()] or ['']
        h += LINE_PT * max(sum(_lines(v, 58) for v in vals), _lines(label, 22))
        if i and kind == 'PARTEI' and not label.lower().startswith('vertreten'):
            party_gaps += 1
            h += 8
        elif i and kind != 'PARTEI':
            value_gaps += 1
            h += 4
    h += 36 + 22  # title
    h += sum(LINE_PT * _lines(t, 90) + 4 for k, t in items if k == 'UNTERTITEL')
    h += sum(22 for k, _ in items if k == 'BEILAGEN')
    free = max(0.0, PAGE_PT - h)
    weights = {'head': 2.0, 'party': 1.0 * party_gaps, 'values': 1.2 if value_gaps else 0.0, 'title': 2.4}
    total = sum(weights.values()) or 1
    caps = {'head': 80, 'party': 30, 'values': 30, 'title': 120}
    out = {}
    for k, w in weights.items():
        share = free * w / total
        out[k] = min(caps[k], share / party_gaps if k == 'party' and party_gaps else share)
    return out


def head(doc, adressat, az, after=10):
    """Addressee top left (court in bold), file number top right."""
    t = doc.add_table(rows=1, cols=2)
    borderless(t)
    widths(t, Cm(10.5), Cm(6.0))
    left, right = t.rows[0].cells
    left.paragraphs[0].text = ''
    lines = adressat or ['An das', '[Gericht]', '[Anschrift]']
    court = 1 if lines[0].lower().startswith('an ') and len(lines) > 1 else 0
    for i, line in enumerate(lines):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        fmt(p, after=0)
        runs(p, line, bold=i == court)
    cell_text(right, az, align=WD_ALIGN_PARAGRAPH.RIGHT)
    fmt(doc.add_paragraph(), after=after)


def parties(doc, rows, party_extra=0, value_extra=0):
    """Rubrum table: label | details (lines split at ';')."""
    t = doc.add_table(rows=0, cols=2)
    borderless(t)
    for i, (kind, label, value) in enumerate(rows):
        r = t.add_row().cells
        lines = [x.strip() for x in value.split(';') if x.strip()] or ['']
        cell_text(r[0], [label])
        cell_text(r[1], lines, bold_first=(kind == 'PARTEI' and not label.lower().startswith('vertreten')))
        new_block = i > 0 and (kind != 'PARTEI' or not label.lower().startswith('vertreten'))
        if new_block:
            first_value = kind != 'PARTEI' and rows[i - 1][0] == 'PARTEI'
            gap = 12 + value_extra if first_value else 12 if kind == 'WEGEN' else 4 if kind != 'PARTEI' else 8 + party_extra
            for c in r:
                c.paragraphs[0].paragraph_format.space_before = Pt(gap)
    widths(t, LABEL_W, VALUE_W)


RUBRUM_KEYS = ('ADRESSAT', 'AZ', 'PARTEI', 'WEGEN', 'STREITWERT', 'RUBRUM')
NUMBERED = re.compile(r'^(\d+(?:\.\d+)+\.?|\d+\.|\(?[a-z]\))\s+(.*)$')  # 1.  1.1  1.1.  a)  (a)


class Body:
    """Text part (page 2 onwards): the first paragraph opens a new page."""

    def __init__(self, doc, new_page=True):
        self.doc, self.pending = doc, new_page

    def par(self):
        p = self.doc.add_paragraph()
        if self.pending:
            p.paragraph_format.page_break_before = True
            self.pending = False
        return p

    def add(self, k, t):
        doc = self.doc
        if k == 'H':
            runs(fmt(self.par(), before=12, after=6, keep_next=True), t, bold=True)
        elif k == 'P':
            p = self.par()
            m = NUMBERED.match(t)
            if m:
                ind = 1.0 if '.' in m.group(1).rstrip('.') else 0.75
                fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left=ind, first=-ind)
                tab_stop(p, ind)
                p.add_run(m.group(1) + '\t')
                runs(p, m.group(2))
            else:
                runs(fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY), t)
        elif k == 'BEWEIS':
            p = fmt(self.par(), after=8, left=3.0, first=-1.75)
            tab_stop(p, 3.0)
            items_ = [x.strip() for x in re.sub(r'^\s*Beweis:\s*', '', t).split(';') if x.strip()]
            p.add_run('Beweis:').bold = True
            p.add_run('\t')
            for i, it in enumerate(items_):
                if i:
                    p.add_run().add_break(WD_BREAK.LINE)
                runs(p, it)
        elif k == 'ANTRAG':
            lead = bool(re.fullmatch(r'[A-ZÄÖÜ ]{4,}[:,]?', t))  # e.g. "URTEILSBEGEHREN:"
            runs(fmt(self.par(), after=6, align=WD_ALIGN_PARAGRAPH.LEFT if lead else WD_ALIGN_PARAGRAPH.JUSTIFY, left=1.5, right=1.5, spacing=1.5, keep_next=lead), t, bold=lead)
        elif k == 'SCHLUSS':
            place, _, who = t.partition('|')
            fmt(self.par(), before=12, after=0)
            tbl = doc.add_table(rows=1, cols=2)
            borderless(tbl)
            widths(tbl, Cm(8.25), Cm(8.25))
            a, b = tbl.rows[0].cells
            cell_text(a, [place.strip()])
            cell_text(b, [x.strip() for x in who.split(';') if x.strip()] or [''], align=WD_ALIGN_PARAGRAPH.RIGHT)


def first_page_title(doc, items, size=14, before=36):
    for k, t in items:
        if k in ('TITEL', 'VERTRAG'):
            p = fmt(doc.add_paragraph(), before=before, after=6, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
            r = p.add_run(t.upper())
            r.bold, r.font.size = True, Pt(size)
            r.font.color.rgb = RGBColor(0, 0, 0)
        elif k == 'UNTERTITEL':
            runs(fmt(doc.add_paragraph(), after=4, align=WD_ALIGN_PARAGRAPH.CENTER), t)
        elif k == 'BEILAGEN':
            label = t if t.lower().startswith('beilage') else 'Beilagen: ' + t
            runs(fmt(doc.add_paragraph(), before=8, after=4, align=WD_ALIGN_PARAGRAPH.CENTER), label, italic=True, size=10)


def render_brief(doc, items):
    adressat = [t for k, t in items if k == 'ADRESSAT']
    az = [t for k, t in items if k == 'AZ' and t]
    rows = []
    for k, t in items:
        label, bar, value = t.partition('|')
        if k == 'PARTEI':
            rows.append(('PARTEI', label.strip(), value.strip()))
        elif k in ('WEGEN', 'STREITWERT', 'RUBRUM'):
            default = {'WEGEN': 'wegen:', 'STREITWERT': 'Streitwert:', 'RUBRUM': ''}[k]
            rows.append((k, label.strip() if bar else default, value.strip() if bar else t))
    sp = spread(adressat, az, rows, items)
    head(doc, adressat, az, after=10 + sp['head'])
    if rows:
        parties(doc, rows, party_extra=sp['party'], value_extra=sp['values'])
    first_page_title(doc, items, before=36 + sp['title'])
    body = Body(doc)
    for k, t in items:
        if k not in RUBRUM_KEYS and k not in ('TITEL', 'UNTERTITEL', 'BEILAGEN'):
            body.add(k, t)


def render_contract(doc, items):
    first_page_title(doc, [(k, t) for k, t in items if k in ('VERTRAG', 'TITEL', 'UNTERTITEL')], size=18, before=140)
    parts = [t for k, t in items if k == 'PARTEI']
    runs(fmt(doc.add_paragraph(), before=36, after=18, align=WD_ALIGN_PARAGRAPH.CENTER), 'abgeschlossen zwischen')
    for i, t in enumerate(parts):
        if i:
            word = 'und' if i == len(parts) - 1 else 'sowie'
            runs(fmt(doc.add_paragraph(), before=6, after=12, align=WD_ALIGN_PARAGRAPH.CENTER), word)
        details, bar, short = t.partition('|')
        lines = [x.strip() for x in details.split(';') if x.strip()] or ['[Partei]']
        for j, line in enumerate(lines):
            runs(fmt(doc.add_paragraph(), after=0, align=WD_ALIGN_PARAGRAPH.CENTER), line, bold=j == 0)
        name = short.strip().strip('„“"')
        if bar and name:
            runs(fmt(doc.add_paragraph(), before=4, after=0, align=WD_ALIGN_PARAGRAPH.CENTER), '(im Folgenden „' + name + '“)', italic=True)
    body = Body(doc)
    signs, annexes = [], []
    for k, t in items:
        if k in ('H', 'P', 'ANTRAG', 'BEWEIS'):
            body.add(k, t)
        elif k == 'UNTERSCHRIFT':
            signs.append(t)
        elif k == 'ANLAGE':
            annexes.append(t)
    if signs:
        fmt(body.par(), before=24, after=0, keep_next=True)
        tbl = doc.add_table(rows=0, cols=2)
        borderless(tbl)
        for i in range(0, len(signs), 2):
            cells = tbl.add_row().cells
            for c, t in zip(cells, signs[i:i + 2]):
                place, _, who = t.partition('|')
                c.paragraphs[0].text = ''
                fmt(c.paragraphs[0], after=36)
                runs(c.paragraphs[0], place.strip() or '[Ort], am [Datum]')
                fmt(c.add_paragraph(), after=0).add_run('_' * 32)
                runs(fmt(c.add_paragraph(), after=18), who.strip() or '[Partei]')
        widths(tbl, Cm(8.25), Cm(8.25))
    if annexes:
        runs(fmt(body.par(), before=18, after=4, keep_next=True), 'Anlagen', bold=True)
        for t in annexes:
            runs(fmt(body.par(), after=0, left=0.5), t)


def render(text, path, title=None):
    """Brief or contract, detected from the markers ([VERTRAG] = contract)."""
    items = parse(text)
    doc = Document()
    setup(doc)
    contract = any(k == 'VERTRAG' for k, _ in items)
    (render_contract if contract else render_brief)(doc, items)
    doc.core_properties.title = title or next((t for k, t in items if k in ('TITEL', 'VERTRAG')), 'Vertrag' if contract else 'Schriftsatz')
    doc.save(path)
    return path


if __name__ == '__main__':
    import sys
    render(open(sys.argv[1], encoding='utf-8').read(), sys.argv[2])
